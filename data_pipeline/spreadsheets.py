"""Export readable temperature workbooks. Inputs: Aligned hourly temperatures. Outputs: Label or prediction XLSX with no analytical worksheets."""

from __future__ import annotations

from datetime import datetime, timedelta
from pathlib import Path
from typing import Iterable
import numpy as np

EXCEL_ROW_LIMIT = 1048576
# Internal non-leap reference for hour offsets, never a weather observation year.
CALENDAR_START = datetime(2001, 1, 1)


def hour_label(hour: int) -> str:
    """Format an annual hour index as month/day and local standard time, without a year."""
    return (CALENDAR_START + timedelta(hours=hour)).strftime("%m-%d %H:%M")


def write_workbook(
    path: Path, sheet_name: str, headers: list[str], rows: Iterable, row_count: int
) -> None:
    """Write a simple streaming workbook and reject Excel row-limit violations."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    if row_count + 1 > EXCEL_ROW_LIMIT:
        raise ValueError(f"{row_count} records exceed the Excel worksheet row limit.")
    if path.exists():
        raise FileExistsError(f"Workbook already exists: {path.name}")
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet(sheet_name)
    sheet.freeze_panes = "A2"
    for index, name in enumerate(headers, 1):
        sheet.column_dimensions[get_column_letter(index)].width = (
            23 if name == "timestamp" else 32 if name == "building_id" else 22
        )
    header_cells = []
    for value in headers:
        cell = WriteOnlyCell(sheet, value=value)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="234E70")
        header_cells.append(cell)
    sheet.append(header_cells)
    observed = 0
    for row in rows:
        cells = []
        for value in row:
            if isinstance(value, np.generic):
                value = value.item()
            cell = WriteOnlyCell(sheet, value=value)
            if isinstance(value, datetime):
                cell.number_format = "mm-dd hh:mm"
            elif isinstance(value, float):
                if not np.isfinite(value):
                    raise ValueError("Workbook contains a non-finite temperature.")
                cell.number_format = "0.0000"
            elif isinstance(value, str):
                cell.data_type = "s"
            cells.append(cell)
        sheet.append(cells)
        observed += 1
    if observed != row_count:
        raise ValueError(f"Expected {row_count} rows, received {observed}.")
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{observed + 1}"
    workbook.save(path)


def outdoor_dry_bulb(epw: Path) -> np.ndarray:
    """Read EPW dry-bulb values in their original hourly row order."""
    with epw.open(encoding="utf-8", errors="strict") as stream:
        lines = list(stream)[8:]
    values = np.asarray(
        [float(line.split(",")[6]) for line in lines if line.strip()], dtype=np.float64
    )
    if values.shape != (8760,) or not np.isfinite(values).all():
        raise ValueError("The demonstration requires exactly 8760 finite hourly EPW records.")
    return values


def export_label(label_path: Path, epw: Path) -> Path:
    """Export NPZ operative temperatures and matching EPW dry-bulb temperatures in a wide sheet."""
    outdoor = outdoor_dry_bulb(epw)
    with np.load(label_path, allow_pickle=False) as label:
        op, zones = label["op"], label["zones"].astype(str)
    if op.shape != (8760, len(zones)):
        raise ValueError(f"Label hours or zones do not match the annual EPW: {label_path.name}")
    rows = ([hour_label(i), outdoor[i], *op[i]] for i in range(8760))
    target = label_path.with_suffix(".xlsx")
    write_workbook(
        target,
        "Temperatures",
        ["timestamp", "outdoor_temperature_C", *[f"op_{zone}_C" for zone in zones]],
        rows,
        len(op),
    )
    return target
